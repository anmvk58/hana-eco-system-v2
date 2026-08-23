from __future__ import annotations

import json
import re
import sys
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


REQUIRED_COLUMNS = ("code", "name", "unit")
OPTIONAL_COLUMNS = (
    "category_id",
    "category_name",
    "sale_price",
    "cost_price",
    "stock_quantity",
    "status",
)
SUPPORTED_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS
FIELD_MAX_LENGTHS = {
    "code": 40,
    "name": 240,
    "unit": 40,
    "category_name": 160,
}
VALID_STATUSES = {"active", "inactive"}


class ApiRequestError(RuntimeError):
    def __init__(self, status_code: int | None, detail: str):
        self.status_code = status_code
        self.detail = detail
        prefix = f"HTTP {status_code}: " if status_code is not None else ""
        super().__init__(prefix + detail)


@dataclass(frozen=True)
class ApiConfig:
    base_url: str
    username: str
    password: str
    timeout_seconds: int = 20


@dataclass(frozen=True)
class ProductImportRow:
    excel_row: int
    code: str
    name: str
    unit: str
    category_id: int | None
    category_name: str | None
    sale_price: Decimal
    cost_price: Decimal
    stock_quantity: Decimal
    status: str


class HanaApiClient:
    def __init__(self, config: ApiConfig):
        self.config = config
        self.access_token: str | None = None

    def _url(self, path: str, query: dict[str, Any] | None = None) -> str:
        url = f"{self.config.base_url.rstrip('/')}/{path.lstrip('/')}"
        if query:
            url += "?" + urlencode(query)
        return url

    def request_json(
        self,
        method: str,
        path: str,
        *,
        payload: dict[str, Any] | None = None,
        query: dict[str, Any] | None = None,
        authenticated: bool = True,
    ) -> Any:
        headers = {"Accept": "application/json"}
        if authenticated:
            if not self.access_token:
                raise ApiRequestError(None, "Chưa đăng nhập API")
            headers["Authorization"] = f"Bearer {self.access_token}"

        body = None
        if payload is not None:
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json; charset=utf-8"

        request = Request(
            self._url(path, query),
            data=body,
            headers=headers,
            method=method,
        )
        try:
            with urlopen(request, timeout=self.config.timeout_seconds) as response:
                response_body = response.read()
                if not response_body:
                    return None
                return json.loads(response_body.decode("utf-8"))
        except HTTPError as exc:
            raw_body = exc.read().decode("utf-8", errors="replace")
            try:
                error_body = json.loads(raw_body)
                detail_value = error_body.get("detail", raw_body)
                detail = (
                    detail_value
                    if isinstance(detail_value, str)
                    else json.dumps(detail_value, ensure_ascii=False)
                )
            except json.JSONDecodeError:
                detail = raw_body or exc.reason
            raise ApiRequestError(exc.code, detail) from exc
        except URLError as exc:
            raise ApiRequestError(
                None,
                f"Không kết nối được API tại {self.config.base_url}: {exc.reason}",
            ) from exc

    def login(self) -> None:
        response = self.request_json(
            "POST",
            "/auth/login",
            payload={
                "username": self.config.username,
                "password": self.config.password,
            },
            authenticated=False,
        )
        token = response.get("access_token") if isinstance(response, dict) else None
        if not token:
            raise ApiRequestError(None, "API đăng nhập thành công nhưng không trả access_token")
        self.access_token = token

    def logout(self) -> None:
        if not self.access_token:
            return
        try:
            self.request_json("POST", "/auth/logout")
        except ApiRequestError as exc:
            print(f"Cảnh báo: không logout được phiên API: {exc}", file=sys.stderr)
        finally:
            self.access_token = None


def normalize_header(value: Any) -> str:
    return str(value or "").strip().lower()


def normalize_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    normalized = str(value).strip()
    return normalized or None


def normalize_code(value: Any) -> str | None:
    normalized = normalize_text(value)
    if normalized is None:
        return None
    return re.sub(r"\s+", "", normalized) or None


def parse_decimal(
    value: Any,
    *,
    field: str,
    excel_row: int,
    default: Decimal = Decimal("0"),
) -> Decimal:
    normalized = normalize_text(value)
    if normalized is None:
        return default
    try:
        parsed = Decimal(normalized.replace(" ", "").replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(
            f"Dòng {excel_row}: cột '{field}' không phải là số hợp lệ"
        ) from exc
    if not parsed.is_finite():
        raise ValueError(f"Dòng {excel_row}: cột '{field}' không phải là số hữu hạn")
    return parsed


def parse_category_id(value: Any, excel_row: int) -> int | None:
    normalized = normalize_text(value)
    if normalized is None:
        return None
    try:
        parsed = Decimal(normalized)
    except InvalidOperation as exc:
        raise ValueError(
            f"Dòng {excel_row}: category_id phải là số nguyên dương"
        ) from exc
    if parsed != parsed.to_integral_value() or parsed <= 0:
        raise ValueError(f"Dòng {excel_row}: category_id phải là số nguyên dương")
    return int(parsed)


def validate_length(field: str, value: str | None, excel_row: int) -> None:
    if value is not None and len(value) > FIELD_MAX_LENGTHS[field]:
        raise ValueError(
            f"Dòng {excel_row}: cột '{field}' vượt quá "
            f"{FIELD_MAX_LENGTHS[field]} ký tự"
        )


def read_product_rows(
    excel_file: Path,
    sheet_name: str | None = None,
) -> list[ProductImportRow]:
    if not excel_file.is_file():
        raise ValueError(f"Không tìm thấy file Excel: {excel_file}")
    if excel_file.suffix.lower() != ".xlsx":
        raise ValueError("Script chỉ hỗ trợ file Excel định dạng .xlsx")

    workbook = load_workbook(excel_file, read_only=True, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(
                    f"Không tìm thấy sheet '{sheet_name}'. "
                    f"Các sheet hiện có: {', '.join(workbook.sheetnames)}"
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active

        row_iterator = worksheet.iter_rows(values_only=True)
        header_row = next(row_iterator, None)
        if header_row is None:
            raise ValueError("File Excel không có dòng tiêu đề")

        header_indexes: dict[str, int] = {}
        for index, value in enumerate(header_row):
            header = normalize_header(value)
            if not header:
                continue
            if header in header_indexes:
                raise ValueError(f"Dòng tiêu đề có cột bị trùng: '{header}'")
            header_indexes[header] = index

        missing_columns = [column for column in REQUIRED_COLUMNS if column not in header_indexes]
        if missing_columns:
            raise ValueError("Thiếu cột bắt buộc: " + ", ".join(missing_columns))

        products: list[ProductImportRow] = []
        seen_codes: dict[str, int] = {}
        for excel_row, values in enumerate(row_iterator, start=2):
            raw_values = {
                column: values[index] if index < len(values) else None
                for column, index in header_indexes.items()
                if column in SUPPORTED_COLUMNS
            }
            if all(normalize_text(raw_values.get(column)) is None for column in SUPPORTED_COLUMNS):
                continue

            code = normalize_code(raw_values.get("code"))
            name = normalize_text(raw_values.get("name"))
            unit = normalize_text(raw_values.get("unit"))
            category_id = parse_category_id(raw_values.get("category_id"), excel_row)
            category_name = normalize_text(raw_values.get("category_name"))
            sale_price = parse_decimal(
                raw_values.get("sale_price"), field="sale_price", excel_row=excel_row
            )
            cost_price = parse_decimal(
                raw_values.get("cost_price"), field="cost_price", excel_row=excel_row
            )
            stock_quantity = parse_decimal(
                raw_values.get("stock_quantity"),
                field="stock_quantity",
                excel_row=excel_row,
            )
            status = (normalize_text(raw_values.get("status")) or "active").lower()

            if not code:
                raise ValueError(f"Dòng {excel_row}: cột 'code' không được để trống")
            if not name:
                raise ValueError(f"Dòng {excel_row}: cột 'name' không được để trống")
            if not unit:
                raise ValueError(f"Dòng {excel_row}: cột 'unit' không được để trống")
            if category_id is not None and category_name is not None:
                raise ValueError(
                    f"Dòng {excel_row}: chỉ nhập category_id hoặc category_name, không nhập cả hai"
                )
            if sale_price < 0:
                raise ValueError(f"Dòng {excel_row}: sale_price không được âm")
            if cost_price < 0:
                raise ValueError(f"Dòng {excel_row}: cost_price không được âm")
            if status not in VALID_STATUSES:
                raise ValueError(
                    f"Dòng {excel_row}: status chỉ nhận 'active' hoặc 'inactive'"
                )

            validate_length("code", code, excel_row)
            validate_length("name", name, excel_row)
            validate_length("unit", unit, excel_row)
            validate_length("category_name", category_name, excel_row)

            code_key = code.casefold()
            if code_key in seen_codes:
                raise ValueError(
                    f"Dòng {excel_row}: code '{code}' trùng với dòng "
                    f"{seen_codes[code_key]} trong Excel"
                )
            seen_codes[code_key] = excel_row

            products.append(
                ProductImportRow(
                    excel_row=excel_row,
                    code=code,
                    name=name,
                    unit=unit,
                    category_id=category_id,
                    category_name=category_name,
                    sale_price=sale_price,
                    cost_price=cost_price,
                    stock_quantity=stock_quantity,
                    status=status,
                )
            )
    finally:
        workbook.close()

    if not products:
        raise ValueError("File Excel không có dòng sản phẩm hợp lệ")
    return products


def resolve_category_ids(
    client: HanaApiClient,
    rows: list[ProductImportRow],
) -> dict[int, int | None]:
    if not any(row.category_id is not None or row.category_name for row in rows):
        return {row.excel_row: None for row in rows}

    categories = client.request_json(
        "GET",
        "/product-categories",
        query={"include_deleted": "false", "limit": 500},
    )
    categories_by_id = {int(category["id"]): category for category in categories}
    categories_by_name = {
        str(category["name"]).strip().casefold(): category for category in categories
    }

    resolved: dict[int, int | None] = {}
    for row in rows:
        if row.category_id is not None:
            if row.category_id not in categories_by_id:
                raise ValueError(
                    f"Dòng {row.excel_row}: không tìm thấy category_id {row.category_id}"
                )
            resolved[row.excel_row] = row.category_id
        elif row.category_name:
            category = categories_by_name.get(row.category_name.casefold())
            if not category:
                raise ValueError(
                    f"Dòng {row.excel_row}: không tìm thấy danh mục '{row.category_name}'"
                )
            resolved[row.excel_row] = int(category["id"])
        else:
            resolved[row.excel_row] = None
    return resolved


def ensure_product_codes_do_not_exist(
    client: HanaApiClient,
    rows: list[ProductImportRow],
) -> None:
    existing_codes: list[str] = []
    for row in rows:
        products = client.request_json(
            "GET",
            "/products",
            query={
                "search": row.code,
                "include_deleted": "true",
                "skip": 0,
                "limit": 200,
            },
        )
        if any(str(product.get("code", "")).casefold() == row.code.casefold() for product in products):
            existing_codes.append(row.code)
    if existing_codes:
        raise ValueError("Mã sản phẩm đã tồn tại: " + ", ".join(existing_codes))


def build_product_payload(
    row: ProductImportRow,
    category_id: int | None,
) -> dict[str, Any]:
    return {
        "code": row.code,
        "name": row.name,
        "category_id": category_id,
        "unit": row.unit,
        "sale_price": str(row.sale_price),
        "cost_price": str(row.cost_price),
        "stock_quantity": str(row.stock_quantity),
        "status": row.status,
    }


def import_products_via_api(
    *,
    excel_file: Path,
    api_config: ApiConfig,
    sheet_name: str | None = None,
    dry_run: bool = False,
) -> int:
    rows = read_product_rows(excel_file.resolve(), sheet_name)
    client = HanaApiClient(api_config)
    client.login()
    try:
        category_ids = resolve_category_ids(client, rows)
        ensure_product_codes_do_not_exist(client, rows)

        if dry_run:
            print(
                f"Kiểm tra thành công {len(rows)} sản phẩm; chưa gọi API tạo sản phẩm."
            )
            return 0

        created_count = 0
        for row in rows:
            payload = build_product_payload(row, category_ids[row.excel_row])
            try:
                created = client.request_json("POST", "/products", payload=payload)
            except ApiRequestError as exc:
                raise ApiRequestError(
                    exc.status_code,
                    f"Dòng {row.excel_row}, code '{row.code}': {exc.detail}. "
                    f"Đã tạo thành công {created_count} sản phẩm trước khi gặp lỗi.",
                ) from exc
            created_count += 1
            print(
                f"[{created_count}/{len(rows)}] Đã tạo {created['code']} - "
                f"{created['name']} (ID {created['id']})"
            )

        print(f"Hoàn tất: đã tạo {created_count} sản phẩm qua API.")
        return created_count
    finally:
        client.logout()


def main() -> int:
    # ================================================================
    # CHỈNH CẤU HÌNH IMPORT TẠI ĐÂY
    # ================================================================
    excel_file = Path(r"C:\Users\Zefus\Downloads\Hana_data\product.xlsx")

    api_config = ApiConfig(
        base_url="http://localhost:8000/api",
        username="admin",
        password="admin",
        timeout_seconds=20,
    )

    sheet_name: str | None = "Sheet2"  # Ví dụ: "SanPham"
    dry_run = False                 # Đổi thành False để tạo sản phẩm thật
    # ================================================================

    try:
        import_products_via_api(
            excel_file=excel_file,
            api_config=api_config,
            sheet_name=sheet_name,
            dry_run=dry_run,
        )
        return 0
    except (ValueError, OSError, ApiRequestError) as exc:
        print(f"Lỗi: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
